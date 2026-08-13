"""
Fix V9: Restore Slide 5 to original backup (no changes).
"""
import os
from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'La finalisima present_backup.pptx')
OUTPUT = os.path.join(BASE, 'La finalisima present_V10.pptx')

FICHE_ICON = os.path.join(BASE, 'fiche_progression_icon.png')
CHECK_ICON = os.path.join(BASE, 'check_icon.png')
PLANIFIER_ICON = os.path.join(BASE, 'planifier_icon.png')
QRCODE_ICON = os.path.join(BASE, 'qrcode_icon.png')
SIGNATURE_ICON = os.path.join(BASE, 'signature_icon.png')
CLOCK_ICON = os.path.join(BASE, 'clock_icon.png')
ICONE_CENTRALISER = os.path.join(BASE, 'icone_centraliser.png')
ICONE_FIABILISER = os.path.join(BASE, 'icone_fiabiliser.png')
ICONE_AUTOMATISER = os.path.join(BASE, 'icone_automatiser.png')
ARROW_ICON = os.path.join(BASE, 'arrow_right.png')

prs = Presentation(INPUT)
TOTAL = len(prs.slides)
print(f"Loaded backup: {TOTAL} slides")

# === SLIDE 1: Keep original, remove BIENVENUE/Soutenance if any ===
slide1 = prs.slides[0]
for shape in list(slide1.shapes):
    if shape.has_text_frame:
        t = shape.text_frame.text
        if 'BIENVENUE' in t or 'Soutenance de Projet' in t:
            shape._element.getparent().remove(shape._element)
            print(f"Slide 1: removed '{t[:60]}'")
for shape in slide1.shapes:
    if shape.name in ('Text 8', 'Text 11'):
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                run.font.bold = True
print("Slide 1: done (original)")

# === SLIDE 2: Replace Apres with Objectifs ===
slide2 = prs.slides[1]
to_del = ['Text 10','Shape 20','Shape 21','Text 22','Text 23',
          'Shape 24','Image 4','Text 25','Text 26',
          'Shape 27','Image 5','Text 28','Text 29',
          'Shape 30','Image 6','Text 31','Text 32']
for shape in list(slide2.shapes):
    if shape.name in to_del:
        shape._element.getparent().remove(shape._element)
print(f"Slide 2: removed {len(to_del)} Apres shapes")

# Objectifs title
tb = slide2.shapes.add_textbox(5500000, 1325880, 3474720, 320040)
p = tb.text_frame.paragraphs[0]
p.text = 'Objectifs'
r = p.runs[0]
r.font.size = Emu(215900)
r.font.bold = True
r.font.color.rgb = RGBColor(0x22, 0xC5, 0x5E)

ITEMS = [
    (1810512, ICONE_CENTRALISER, 'Centraliser', 'Tout au meme endroit'),
    (2926080, ICONE_FIABILISER,  'Fiabiliser',   'Donnees verifiees'),
    (4041648, ICONE_AUTOMATISER, 'Automatiser',  'Calculs et rapports'),
]
for y, icon, title, sub in ITEMS:
    box = slide2.shapes.add_shape(1, 5500000, y, 3474720, 914400)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x17, 0x20, 0x33)
    box.line.color.rgb = RGBColor(0x10, 0xB9, 0x81)
    box.line.width = Emu(14605)
    slide2.shapes.add_picture(icon, 5701168, y+45720, 786384, 786384)
    tb1 = slide2.shapes.add_textbox(6643000, y+237744, 2100000, 256032)
    pt = tb1.text_frame.paragraphs[0]
    pt.text = title; rt = pt.runs[0]
    rt.font.size = Emu(215900); rt.font.bold = True
    rt.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tb2 = slide2.shapes.add_textbox(6643000, y+565000, 2100000, 219456)
    ps = tb2.text_frame.paragraphs[0]
    ps.text = sub; rs = ps.runs[0]
    rs.font.size = Emu(158750)
    rs.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

slide2.shapes.add_picture(ARROW_ICON, 4300000, 3000000, 800000, 500000)
print("Slide 2: Objectifs + arrow done")

# === SLIDE 3: Fiche progression icon ===
slide3 = prs.slides[2]
for shape in list(slide3.shapes):
    if shape.name == 'Image 4':
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
        shape._element.getparent().remove(shape._element)
        slide3.shapes.add_picture(FICHE_ICON, l, t, w, h)
        print("Slide 3: fiche progression icon replaced")
        break

# === SLIDE 4: Replace icons (Planifier, QR Code, Valider) ===
slide4 = prs.slides[3]
for shape in list(slide4.shapes):
    if shape.name == 'Image 4' and shape.left > 8000000:
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
        shape._element.getparent().remove(shape._element)
        slide4.shapes.add_picture(CHECK_ICON, l, t, w, h)
        print("Slide 4: Valider -> check")
    elif shape.name == 'Image 1':
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
        shape._element.getparent().remove(shape._element)
        slide4.shapes.add_picture(PLANIFIER_ICON, l, t, w, h)
        print("Slide 4: Planifier -> calendar")
    elif shape.name == 'Image 2':
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
        shape._element.getparent().remove(shape._element)
        slide4.shapes.add_picture(QRCODE_ICON, l, t, w, h)
        print("Slide 4: QR Code -> QR")

# === SLIDE 5: NO CHANGES - KEEP ORIGINAL ===
print("Slide 5: KEPT ORIGINAL (no changes)")

# === SLIDE 6: NO CHANGES ===
print("Slide 6: no changes")

# === SLIDE 7: NO CHANGES ===
print("Slide 7: no changes")

# === SLIDE 8: Clock + Signature icons ===
slide8 = prs.slides[7]
for shape in list(slide8.shapes):
    if shape.name == 'Image 3':
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
        shape._element.getparent().remove(shape._element)
        slide8.shapes.add_picture(CLOCK_ICON, l, t, w, h)
        print("Slide 8: Honoraires -> clock")
    if shape.name == 'Image 7':
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
        shape._element.getparent().remove(shape._element)
        slide8.shapes.add_picture(SIGNATURE_ICON, l, t, w, h)
        print("Slide 8: Signature -> signature icon")

# === SLIDE 9: Remove Questions if present ===
slide9 = prs.slides[8]
for shape in list(slide9.shapes):
    if shape.has_text_frame and 'Questions' in shape.text_frame.text:
        shape._element.getparent().remove(shape._element)
        print("Slide 9: Questions removed")
        break
print("Slide 9: MERCI done")

# === UPDATE SLIDE NUMBERS ===
for i, slide in enumerate(prs.slides):
    sn = i + 1
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if '/' in t and len(t) <= 8:
                parts = t.split('/')
                if len(parts) == 2 and parts[0].strip().isdigit():
                    new_t = f'{sn} / {TOTAL}'
                    for p in shape.text_frame.paragraphs:
                        if p.text.strip() != new_t:
                            for run in p.runs:
                                run.text = ''
                            if p.runs:
                                p.runs[0].text = new_t
print(f"Slide numbers -> X/{TOTAL}")

prs.save(OUTPUT)
print(f"\nDONE: {OUTPUT}")
print(f"Total slides: {len(prs.slides)}")
