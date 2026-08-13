"""Read slide 8 (Perspectives) from temp V3 copy."""
import os
from pptx import Presentation

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'temp_v3_read.pptx')

prs = Presentation(INPUT)
print(f"Total slides: {len(prs.slides)}")

for idx in [6, 7, 8]:
    slide = prs.slides[idx]
    print(f"\n=== Slide {idx+1} (index {idx}) ===")
    for i, shape in enumerate(slide.shapes):
        stype = str(shape.shape_type)
        text = ""
        if shape.has_text_frame:
            text = shape.text_frame.text[:150]
        print(f"[{i}] name='{shape.name}' type={stype} pos=({shape.left},{shape.top}) size=({shape.width},{shape.height}) text='{text}'")
