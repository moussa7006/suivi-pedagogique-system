"""Check slide 5 shapes in temp_v7."""
import os
from pptx import Presentation

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'temp_v7_read.pptx')

prs = Presentation(INPUT)

print("=== Slide 5 (index 4) ===")
slide = prs.slides[4]
for i, shape in enumerate(slide.shapes):
    stype = str(shape.shape_type)
    text = ""
    if shape.has_text_frame:
        text = shape.text_frame.text[:100]
    print(f"[{i}] name='{shape.name}' type={stype} pos=({shape.left},{shape.top}) size=({shape.width},{shape.height}) text='{text}'")

print("\n=== Slide 4 (index 3) ===")
slide4 = prs.slides[3]
for i, shape in enumerate(slide4.shapes):
    stype = str(shape.shape_type)
    text = ""
    if shape.has_text_frame:
        text = shape.text_frame.text[:100]
    print(f"[{i}] name='{shape.name}' type={stype} pos=({shape.left},{shape.top}) size=({shape.width},{shape.height}) text='{text}'")
