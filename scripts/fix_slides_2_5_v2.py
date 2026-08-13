"""Modify slides 2 and 5 on V4 (which has clock icon): Apres+arrow, MVC arch."""
import os
from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'temp_v4_read.pptx')
OUTPUT = os.path.join(BASE, 'La finalisima present_V5.pptx')

ARROW = os.path.join(BASE, 'arrow_right.png')
MVC_ARCH = os.path.join(BASE, 'architecture_mvc.png')

prs = Presentation(INPUT)

# ============================
# SLIDE 2 (index 1)
# ============================
slide2 = prs.slides[1]

# 1. Change "Objectifs" to "Après" in GREEN (#22C55E)
for shape in slide2.shapes:
    if shape.name == 'TextBox 25':
        p = shape.text_frame.paragraphs[0]
        p.text = 'Après'
        for run in p.runs:
            run.font.color.rgb = RGBColor(0x22, 0xC5, 0x5E)
        print("Slide 2: 'Objectifs' -> 'Après' (green #22C55E)")
        break

# 2. Add green arrow between Avant and Apres
slide2.shapes.add_picture(ARROW, 4300000, 3000000, 800000, 500000)
print("Slide 2: Arrow added at center")

# ============================
# SLIDE 5 (index 4)
# ============================
slide5 = prs.slides[4]

# Remove old architecture image (Image 1)
for shape in list(slide5.shapes):
    if shape.name == 'Image 1':
        shape._element.getparent().remove(shape._element)
        print("Slide 5: Old Image 1 removed")
        break

# Add MVC architecture
slide5.shapes.add_picture(MVC_ARCH, 1200000, 1100000, 9800000, 4400000)
print("Slide 5: MVC architecture added")

prs.save(OUTPUT)
print(f"\nSUCCESS -> {OUTPUT}")
print("(Preserves all V4 changes: clock icon on slide 8, etc.)")
