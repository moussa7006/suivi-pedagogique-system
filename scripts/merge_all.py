"""Apply slide 2+5 changes to V6 -> V7 FINAL."""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'temp_v6_read.pptx')
OUTPUT = os.path.join(BASE, 'La finalisima present_V7.pptx')
ARROW = os.path.join(BASE, 'arrow_right.png')
MVC_ARCH = os.path.join(BASE, 'architecture_mvc.png')

prs = Presentation(INPUT)

# === SLIDE 2: "Objectifs" -> "Après" green + arrow ===
slide2 = prs.slides[1]
for shape in slide2.shapes:
    if shape.name == 'TextBox 25':
        p = shape.text_frame.paragraphs[0]
        p.text = 'Après'
        for run in p.runs:
            run.font.color.rgb = RGBColor(0x22, 0xC5, 0x5E)
        print("Slide 2: 'Objectifs' -> 'Après' (green)")
        break
slide2.shapes.add_picture(ARROW, 4300000, 3000000, 800000, 500000)
print("Slide 2: Arrow added")

# === SLIDE 5: Replace with MVC architecture ===
slide5 = prs.slides[4]
for shape in list(slide5.shapes):
    if shape.name == 'Image 1':
        shape._element.getparent().remove(shape._element)
        print("Slide 5: Old Image 1 removed")
        break
slide5.shapes.add_picture(MVC_ARCH, 1200000, 1100000, 9800000, 4400000)
print("Slide 5: MVC architecture added")

prs.save(OUTPUT)
print(f"\n✅ FINAL -> {OUTPUT}")
print("Includes: Slide1(centered)+Slide2(Après+arrow)+Slide5(MVC)+Slide8(clock)")
