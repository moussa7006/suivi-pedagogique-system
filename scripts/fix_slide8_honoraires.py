"""Replace honoraires icon (Image 3) with clock icon on slide 8 of V3."""
import os
from pptx import Presentation
from pptx.util import Emu

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'temp_v3_read.pptx')
OUTPUT = os.path.join(BASE, 'La finalisima present_V4.pptx')
CLOCK = os.path.join(BASE, 'clock_icon.png')

prs = Presentation(INPUT)

# Slide 8 = index 7 (Conclusion & perspectives)
slide = prs.slides[7]

# Find Image 3 (the honoraires icon)
target = None
for shape in slide.shapes:
    if shape.name == 'Image 3':
        target = shape
        break

if target is None:
    print("Image 3 not found!")
else:
    print(f"Found Image 3 at pos=({target.left},{target.top}) size=({target.width},{target.height})")
    
    # Store position and size
    left = target.left
    top = target.top
    width = target.width
    height = target.height
    
    # Remove old image
    target._element.getparent().remove(target._element)
    print("Old Image 3 removed.")
    
    # Add new clock icon at same position
    slide.shapes.add_picture(CLOCK, left, top, width, height)
    print("Clock icon added.")
    
    prs.save(OUTPUT)
    print(f"\nSUCCESS → Saved to {OUTPUT}")
