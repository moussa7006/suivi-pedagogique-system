"""Modify slides 2 and 5: Apres+arrow, MVC architecture."""
import os
from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'temp_v3_read.pptx')
OUTPUT = os.path.join(BASE, 'La finalisima present_V5.pptx')

ARROW = os.path.join(BASE, 'arrow_right.png')
MVC_ARCH = os.path.join(BASE, 'architecture_mvc.png')

prs = Presentation(INPUT)

# ============================
# SLIDE 2 (index 1)
# ============================
slide2 = prs.slides[1]

# 1. Change "Objectifs" title to "Après" in GREEN
for shape in slide2.shapes:
    if shape.name == 'TextBox 25':
        p = shape.text_frame.paragraphs[0]
        p.text = 'Après'
        for run in p.runs:
            run.font.color.rgb = RGBColor(0x22, 0xC5, 0x5E)  # #22C55E green
        print(f"Slide 2: 'Objectifs' -> 'Après' (green)")
        break

# 2. Add green arrow between Avant and Apres
# The gap between columns: left ends ~4,133,088, right starts ~5,500,000
# Center gap: ~4,800,000
# We want arrow centered vertically (items start at y=1,800,000)
arrow_x = 4300000
arrow_y = 3000000
arrow_w = 800000
arrow_h = 500000
slide2.shapes.add_picture(ARROW, arrow_x, arrow_y, arrow_w, arrow_h)
print(f"Slide 2: Arrow added at ({arrow_x},{arrow_y})")

# ============================
# SLIDE 5 (index 4)
# ============================
slide5 = prs.slides[4]

# Remove old architecture image (Image 1)
for shape in list(slide5.shapes):
    if shape.name == 'Image 1':
        shape._element.getparent().remove(shape._element)
        print(f"Slide 5: Old Image 1 removed")
        break

# Add new MVC architecture image
# Position it nicely on the slide
mvc_x = 1200000
mvc_y = 1100000
mvc_w = 9800000
mvc_h = 4400000
slide5.shapes.add_picture(MVC_ARCH, mvc_x, mvc_y, mvc_w, mvc_h)
print(f"Slide 5: MVC architecture added")

# Save
prs.save(OUTPUT)
print(f"\nSUCCESS -> {OUTPUT}")
