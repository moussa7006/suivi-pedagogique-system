"""Analyze slide 4 (and 5,6,7) from backup."""
import os
from pptx import Presentation

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'La finalisima present_backup.pptx')

prs = Presentation(INPUT)

for idx in [3, 4, 5, 6]:
    slide = prs.slides[idx]
    print(f"\n{'='*60}")
    print(f"=== Slide {idx+1} (index {idx}) ===")
    print(f"{'='*60}")
    for i, shape in enumerate(slide.shapes):
        stype = str(shape.shape_type)
        text = ""
        if shape.has_text_frame:
            text = shape.text_frame.text[:200]
        print(f"[{i}] name='{shape.name}' type={stype} pos=({shape.left},{shape.top}) size=({shape.width},{shape.height}) text='{text}'")
