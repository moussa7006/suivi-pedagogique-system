"""Quick fix: replace slide 5 architecture and save as V8."""
import os
from pptx import Presentation

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'temp_v7_read.pptx')
OUTPUT = os.path.join(BASE, 'La finalisima present_V8.pptx')
ARCH = os.path.join(BASE, 'architecture_pro.png')

prs = Presentation(INPUT)
slide = prs.slides[4]

# Find any picture on slide 5 at architecture position and replace
for shape in list(slide.shapes):
    if shape.shape_type == 13 and shape.width > 5000000:  # PICTURE, large = architecture
        left, top, w, h = shape.left, shape.top, shape.width, shape.height
        print(f"Removing '{shape.name}' at ({left},{top}) size=({w},{h})")
        shape._element.getparent().remove(shape._element)
        slide.shapes.add_picture(ARCH, left, top, w, h)
        print("Architecture replaced!")
        break

prs.save(OUTPUT)
print(f"✅ -> {OUTPUT}")
