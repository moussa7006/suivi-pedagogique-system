"""Analyze slides 2 and 5 from temp copy."""
import os
from pptx import Presentation

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'temp_v3_read.pptx')

prs = Presentation(INPUT)

for idx in [1, 4]:  # slide 2 = index 1, slide 5 = index 4
    slide = prs.slides[idx]
    print(f"\n{'='*60}")
    print(f"=== Slide {idx+1} (index {idx}) ===")
    print(f"{'='*60}")
    for i, shape in enumerate(slide.shapes):
        stype = str(shape.shape_type)
        text = ""
        if shape.has_text_frame:
            text = shape.text_frame.text[:180]
        print(f"[{i}] name='{shape.name}' type={stype} pos=({shape.left},{shape.top}) size=({shape.width},{shape.height}) text='{text}'")
