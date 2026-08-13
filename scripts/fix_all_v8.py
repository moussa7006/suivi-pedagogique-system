"""Fix both slide 4 icons AND slide 5 architecture -> V8 FINAL."""
import os
from pptx import Presentation

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'temp_v7_read.pptx')
OUTPUT = os.path.join(BASE, 'La finalisima present_V8.pptx')
PLANIFIER = os.path.join(BASE, 'planifier_icon.png')
QRCODE = os.path.join(BASE, 'qrcode_icon.png')
ARCH = os.path.join(BASE, 'architecture_pro.png')

prs = Presentation(INPUT)

# === SLIDE 4 ===
slide4 = prs.slides[3]
for shape in list(slide4.shapes):
    if shape.name == 'Image 1':
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
        shape._element.getparent().remove(shape._element)
        slide4.shapes.add_picture(PLANIFIER, l, t, w, h)
        print("Slide 4: Planifier -> calendrier")
    elif shape.name == 'Image 2':
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
        shape._element.getparent().remove(shape._element)
        slide4.shapes.add_picture(QRCODE, l, t, w, h)
        print("Slide 4: QR Code -> QR icon")

# === SLIDE 5 ===
slide5 = prs.slides[4]
for shape in list(slide5.shapes):
    if shape.shape_type == 13 and shape.width > 5000000:
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
        shape._element.getparent().remove(shape._element)
        slide5.shapes.add_picture(ARCH, l, t, w, h)
        print("Slide 5: Architecture -> PRO visible version")

prs.save(OUTPUT)
print(f"\n✅ FINAL -> {OUTPUT}")
