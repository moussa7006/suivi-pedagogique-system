"""Fix slide 2: replace "Après" section with "Objectifs" items."""
import os
from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'La finalisima present_V2.pptx')
OUTPUT = os.path.join(BASE, 'La finalisima present_V3.pptx')

prs = Presentation(INPUT)
slide = prs.slides[1]

# ── 1. Supprimer centre + droite ──
to_delete = [
    'Text 10', 'Shape 20', 'Shape 21', 'Text 22', 'Text 23',
    'Shape 24', 'Image 4', 'Text 25', 'Text 26',
    'Shape 27', 'Image 5', 'Text 28', 'Text 29',
    'Shape 30', 'Image 6', 'Text 31', 'Text 32'
]

for shape in list(slide.shapes):
    if shape.name in to_delete:
        shape._element.getparent().remove(shape._element)

print(f"Supprimé {len(to_delete)} éléments.")

# ── 2. Titre "Objectifs" ──
title_box = slide.shapes.add_textbox(5500000, 1325880, 3474720, 320040)
p = title_box.text_frame.paragraphs[0]
p.text = 'Objectifs'
run = p.runs[0]
run.font.size = Emu(215900)
run.font.bold = True
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# ── 3. 3 items Objectifs ──
ITEMS = [
    {'y': 1810512, 'icon': 'icone_centraliser.png', 'title': 'Centraliser', 'sub': 'Tout au même endroit'},
    {'y': 2926080, 'icon': 'icone_fiabiliser.png', 'title': 'Fiabiliser', 'sub': 'Données vérifiées'},
    {'y': 4041648, 'icon': 'icone_automatiser.png', 'title': 'Automatiser', 'sub': 'Calculs et rapports'},
]

BOX_X = 5500000
BOX_W = 3474720
BOX_H = 914400

for item in ITEMS:
    y = item['y']

    # Boîte arrondie
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, BOX_X, y, BOX_W, BOX_H)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x17, 0x20, 0x33)
    box.line.color.rgb = RGBColor(0x10, 0xB9, 0x81)
    box.line.width = Emu(14605)

    # Icône
    icon_path = os.path.join(BASE, item['icon'])
    slide.shapes.add_picture(icon_path, BOX_X + 201168, y + 45720, 786384, 786384)

    # Titre
    tb = slide.shapes.add_textbox(BOX_X + 1143000, y + 237744, 2100000, 256032)
    p_t = tb.text_frame.paragraphs[0]
    p_t.text = item['title']
    r_t = p_t.runs[0]
    r_t.font.size = Emu(215900)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Sous-titre
    tb2 = slide.shapes.add_textbox(BOX_X + 1143000, y + 565000, 2100000, 219456)
    p_s = tb2.text_frame.paragraphs[0]
    p_s.text = item['sub']
    r_s = p_s.runs[0]
    r_s.font.size = Emu(158750)
    r_s.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

prs.save(OUTPUT)
print(f"SUCCESS → {OUTPUT}")
