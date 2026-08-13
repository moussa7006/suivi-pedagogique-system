"""Analyze slide 1 from V4 copy."""
import os
from pptx import Presentation

BASE = r'C:\Users\boura\OneDrive\Documents\suivi-pedagogique-system\MEMOIRE'
INPUT = os.path.join(BASE, 'temp_v4_read.pptx')

prs = Presentation(INPUT)
slide = prs.slides[0]

print(f"=== Slide 1 (index 0) ===\n")
for i, shape in enumerate(slide.shapes):
    stype = str(shape.shape_type)
    text = ""
    if shape.has_text_frame:
        text = shape.text_frame.text[:200]
    print(f"[{i}] name='{shape.name}' type={stype}")
    print(f"    pos=({shape.left},{shape.top}) size=({shape.width},{shape.height})")
    print(f"    text='{text}'")
    # Check fill color if it's a shape
    if hasattr(shape, 'fill'):
        try:
            fill_type = str(shape.fill.type)
            print(f"    fill_type={fill_type}")
        except:
            pass
    print()
