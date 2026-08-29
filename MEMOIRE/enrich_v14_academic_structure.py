from copy import deepcopy
from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SOURCE = Path('La finalisima present_V14.before-academic-structure.pptx')
OUTPUT = Path('La finalisima present_V15.pptx')
LOGO = Path('_edutrack_logo.png')

NAVY = RGBColor(15, 66, 114)
BLUE = RGBColor(24, 104, 171)
SKY = RGBColor(103, 181, 229)
PALE = RGBColor(232, 245, 253)
BACKGROUND = RGBColor(248, 250, 252)
WHITE = RGBColor(255, 255, 255)
SLATE = RGBColor(48, 65, 82)
MUTED = RGBColor(101, 122, 140)
BORDER = RGBColor(211, 226, 236)
GREEN = RGBColor(16, 135, 105)
ORANGE = RGBColor(214, 121, 32)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color


def add_text(slide, text, x, y, w, h, size, color=SLATE, bold=False,
             align=PP_ALIGN.LEFT, font='Aptos', valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = text
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_rect(slide, x, y, w, h, color, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(shape, color)
    if rounded:
        shape.adjustments[0] = 0.10
    return shape


def add_circle(slide, x, y, d, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    fill(shape, color)
    return shape


def add_header(slide, title, subtitle):
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = BACKGROUND
    add_rect(slide, 0, 0, 13.333, 0.12, BLUE)
    add_rect(slide, 0.42, 0.48, 0.72, 0.06, BLUE)
    add_text(slide, title, 0.42, 0.68, 9.4, 0.56, 27, NAVY, True, font='Aptos Display')
    add_text(slide, subtitle, 0.42, 1.28, 8.8, 0.30, 11.5, MUTED)
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(10.88), Inches(0.47), width=Inches(1.55), height=Inches(0.45))
    add_rect(slide, 0.42, 7.12, 1.48, 0.045, SKY)


def add_card(slide, x, y, w, h, index, title, body, accent=BLUE):
    card = add_rect(slide, x, y, w, h, WHITE, rounded=True)
    card.line.color.rgb = BORDER
    card.line.width = Pt(1)
    add_circle(slide, x + 0.28, y + 0.25, 0.46, accent)
    add_text(slide, index, x + 0.28, y + 0.25, 0.46, 0.46, 11, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, title, x + 0.92, y + 0.23, w - 1.18, 0.36, 15, NAVY, True)
    add_text(slide, body, x + 0.30, y + 0.78, w - 0.60, h - 0.96, 11.5, SLATE, False, valign=MSO_ANCHOR.TOP)


def add_new_slide(prs, title, subtitle, cards):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    add_header(slide, title, subtitle)
    count = len(cards)
    if count == 3:
        positions = [(0.42, 2.00, 4.02, 3.94), (4.66, 2.00, 4.02, 3.94), (8.90, 2.00, 4.02, 3.94)]
    elif count == 4:
        positions = [(0.42, 1.92, 6.10, 2.10), (6.80, 1.92, 6.10, 2.10), (0.42, 4.30, 6.10, 2.10), (6.80, 4.30, 6.10, 2.10)]
    else:
        positions = [(0.42, 2.00, 12.48, 3.90)]
    accents = [BLUE, GREEN, ORANGE, SKY]
    for i, ((title_text, body), (x, y, w, h)) in enumerate(zip(cards, positions), 1):
        add_card(slide, x, y, w, h, str(i), title_text, body, accents[(i - 1) % len(accents)])
    return slide


def replace_exact_text(slide, replacements):
    for shape in slide.shapes:
        if not getattr(shape, 'has_text_frame', False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text in replacements:
                    run.text = replacements[run.text]


def reorder_slides(prs, ordered_slides):
    sld_id_list = prs.slides._sldIdLst
    ids_by_part = {slide.part: element for slide, element in zip(prs.slides, sld_id_list)}
    ordered_elements = [ids_by_part[slide.part] for slide in ordered_slides]
    for element in list(sld_id_list):
        sld_id_list.remove(element)
    for element in ordered_elements:
        sld_id_list.append(element)


def update_numbers(prs):
    total = len(prs.slides)
    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not getattr(shape, 'has_text_frame', False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    normalized = run.text.strip()
                    if re.fullmatch(r'\d{1,2}\s*/\s*\d{1,2}', normalized):
                        run.text = f'{index:02d} / {total:02d}'


prs = Presentation(SOURCE)
original = list(prs.slides)

# Existing slide 2 becomes the dedicated problem statement slide.
replace_exact_text(original[1], {
    'Problématique & objectifs': 'Problématique',
    'Objectifs': 'Conséquences observées',
    'Centraliser': 'Information fragmentée',
    'Tout au même endroit': 'Décisions ralenties',
    'Fiabiliser': 'Erreurs de suivi',
    'Données vérifiées': 'Présences difficiles à prouver',
    'Automatiser': 'Calculs manuels',
    'Calculs et rapports': 'Honoraires et rapports peu fiables',
})

introduction = add_new_slide(
    prs,
    'Introduction : pourquoi EduTrack ?',
    'Contexte du projet — moderniser le suivi pédagogique au sein d’un établissement',
    [
        ('Un suivi au cœur de la qualité', 'La planification des cours, les présences et les contenus enseignés doivent rester accessibles, cohérents et vérifiables.'),
        ('Des acteurs à coordonner', 'Administrateurs et enseignants ont besoin d’informations fiables, partagées au bon moment et adaptées à leur rôle.'),
        ('Notre réponse', 'EduTrack réunit une application web d’administration et une application mobile pour les enseignants.'),
    ],
)

objectives = add_new_slide(
    prs,
    'Objectifs & hypothèses de travail',
    'Ce que la solution doit améliorer et les effets attendus',
    [
        ('Objectif général', 'Concevoir une solution web et mobile pour centraliser et fiabiliser le suivi pédagogique.'),
        ('Objectifs spécifiques', 'Planifier les séances, suivre les émargements, documenter la progression et automatiser les honoraires.'),
        ('Hypothèse H1', 'La centralisation des données améliore la disponibilité de l’information et le pilotage pédagogique.'),
        ('Hypothèses H2 à H4', 'Le QR code renforce la traçabilité ; le mobile simplifie l’émargement ; les séances validées fiabilisent les honoraires.'),
    ],
)

review = add_new_slide(
    prs,
    'Résumé de la revue critique',
    'Analyse des pratiques existantes et justification de la solution proposée',
    [
        ('Pratiques actuelles', 'Excel et formulaires papier restent utiles mais dispersent les données entre plusieurs supports.'),
        ('Limites identifiées', 'Double saisie, recherche lente, manque d’historique consolidé et difficulté de contrôle des présences.'),
        ('Apport d’EduTrack', 'Un parcours unique relie la séance, le QR code, l’émargement, la fiche de progression et les honoraires.'),
    ],
)

methodology = add_new_slide(
    prs,
    'Cadre méthodologique & analytique',
    'Démarche suivie pour concevoir, réaliser et vérifier la solution',
    [
        ('Analyse des besoins', 'Identification des acteurs, des processus existants et des contraintes de gestion pédagogique.'),
        ('Conception', 'Modélisation UML : cas d’utilisation, séquence d’émargement et architecture MVC.'),
        ('Réalisation', 'API REST Spring Boot sécurisée par JWT, interface Angular/Ionic et base PostgreSQL.'),
        ('Validation', 'Tests backend, build frontend et démonstration des parcours administrateur et enseignant.'),
    ],
)

results = add_new_slide(
    prs,
    'Résultats & discussion',
    'Fonctionnalités réalisées et bénéfices observés dans les parcours métier',
    [
        ('Résultats fonctionnels', 'Gestion des référentiels, emplois du temps, séances, QR code, émargements, fiches de progression et honoraires.'),
        ('Résultats techniques', 'Application web et mobile, authentification JWT, contrôle des rôles, API documentée et données centralisées.'),
        ('Discussion', 'Le lien séance → émargement → fiche → honoraire réduit les ruptures de suivi et améliore la traçabilité.'),
    ],
)

# Slide order: cover; new introduction; existing problem; new academic slides;
# then the original product walkthrough; new results before conclusion and closure.
ordered = [
    original[0], introduction, original[1], objectives, review, methodology,
    original[2], original[3], original[4], original[5], original[6],
    results, original[7], original[8],
]
reorder_slides(prs, ordered)
update_numbers(prs)
prs.save(OUTPUT)
print(f'Academic structure applied: {OUTPUT} ({len(prs.slides)} slides)')
