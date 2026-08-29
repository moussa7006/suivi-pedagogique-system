from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE_TYPE

p = Presentation('La finalisima present_V11.pptx')
for i, slide in enumerate(p.slides, 1):
    print(f'--- slide {i} ---')
    for j, sh in enumerate(slide.shapes):
        kind = str(sh.shape_type)
        text = getattr(sh, 'text', '').replace('\n', ' / ')[:80]
        fill = ''
        try:
            if sh.fill.type is not None:
                fill = f' fill={sh.fill.type} rgb={getattr(sh.fill.fore_color, "rgb", None)}'
        except Exception:
            pass
        print(j, kind, f'x={sh.left/914400:.2f} y={sh.top/914400:.2f} w={sh.width/914400:.2f} h={sh.height/914400:.2f}', fill, repr(text))
