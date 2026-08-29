from pptx import Presentation
p = Presentation('La finalisima present_V11.pptx')
for n, s in enumerate(p.slides, 1):
    print(f'--- SLIDE {n} ---')
    for x in s.shapes:
        if getattr(x, 'has_text_frame', False) and x.text.strip():
            print(x.text.replace('\n', ' / '))
