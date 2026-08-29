from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SOURCE = Path('La finalisima present_V18.before-conclusion-merge.pptx')
OUTPUT = Path('La finalisima present_V19.pptx')
LOGO = Path('_edutrack_logo.png')
NAVY, BLUE, SKY = RGBColor(15, 66, 114), RGBColor(24, 104, 171), RGBColor(103, 181, 229)
BACKGROUND, WHITE = RGBColor(248, 250, 252), RGBColor(255, 255, 255)
SLATE, MUTED, BORDER = RGBColor(48, 65, 82), RGBColor(101, 122, 140), RGBColor(211, 226, 236)
GREEN, ORANGE = RGBColor(16, 135, 105), RGBColor(214, 121, 32)


def fill(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.color.rgb = color


def add_text(slide, value, x, y, w, h, size, color=SLATE, bold=False, align=PP_ALIGN.LEFT, font='Aptos', valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame; frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]; paragraph.alignment = align; paragraph.text = value
    for run in paragraph.runs:
        run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, color, rounded=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(shape, color)
    if rounded: shape.adjustments[0] = 0.10
    return shape


def circle(slide, x, y, d, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)); fill(shape, color); return shape


def clear(slide):
    tree = slide.shapes._spTree
    for shape in list(slide.shapes): tree.remove(shape._element)


def header(slide, title, subtitle):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BACKGROUND
    rect(slide, 0, 0, 13.333, 0.12, BLUE); rect(slide, 0.42, 0.48, 0.72, 0.06, BLUE)
    add_text(slide, title, 0.42, 0.68, 10, 0.56, 27, NAVY, True, font='Aptos Display')
    add_text(slide, subtitle, 0.42, 1.28, 9.7, 0.30, 11.5, MUTED)
    if LOGO.exists(): slide.shapes.add_picture(str(LOGO), Inches(10.88), Inches(0.47), width=Inches(1.55), height=Inches(0.45))
    rect(slide, 0.42, 7.12, 1.48, 0.045, SKY)


def card(slide, x, y, w, h, number, title, body, accent):
    shape = rect(slide, x, y, w, h, WHITE, True); shape.line.color.rgb = BORDER; shape.line.width = Pt(1)
    circle(slide, x + 0.28, y + 0.25, 0.46, accent)
    add_text(slide, str(number), x + 0.28, y + 0.25, 0.46, 0.46, 11, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, title, x + 0.92, y + 0.23, w - 1.18, 0.36, 15, NAVY, True)
    add_text(slide, body, x + 0.30, y + 0.78, w - 0.60, h - 0.96, 11.5, SLATE, False, valign=MSO_ANCHOR.TOP)


def reorder(prs, slides):
    mapping = {slide.part: element for slide, element in zip(prs.slides, prs.slides._sldIdLst)}
    elements = [mapping[slide.part] for slide in slides]
    for element in list(prs.slides._sldIdLst): prs.slides._sldIdLst.remove(element)
    for element in elements: prs.slides._sldIdLst.append(element)


def renumber(prs):
    total = len(prs.slides)
    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not getattr(shape, 'has_text_frame', False): continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if re.fullmatch(r'\d{1,2}\s*/\s*\d{1,2}', run.text.strip()): run.text = f'{index:02d} / {total:02d}'

prs = Presentation(SOURCE)
slides = list(prs.slides)
intro, conclusion, recommendations = slides[1], slides[11], slides[12]

# Intro: a deliberately brief contextual entry, without explaining features too early.
clear(intro); header(intro, 'Introduction', 'Contexte général du projet')
add_text(intro, 'Dans les établissements d’enseignement, le suivi pédagogique repose souvent\nsur plusieurs supports : documents papier, tableaux Excel et échanges informels.', 1.00, 2.03, 11.35, 0.90, 20, NAVY, True, align=PP_ALIGN.CENTER, font='Aptos Display')
add_text(intro, 'Cette organisation rend la centralisation des informations, la traçabilité des activités\net le suivi des enseignants plus difficiles.', 1.20, 3.38, 10.95, 0.62, 16, SLATE, align=PP_ALIGN.CENTER)
rect(intro, 2.15, 4.55, 9.03, 0.82, WHITE, True)
add_text(intro, 'C’est dans ce contexte que s’inscrit le projet EduTrack.', 2.43, 4.78, 8.48, 0.30, 16, BLUE, True, align=PP_ALIGN.CENTER)

# One final academic slide: conclusion first, then recommendations and perspectives.
clear(conclusion); header(conclusion, 'Conclusion, recommandations & perspectives', 'Bilan du projet et pistes d’amélioration')
items = [
    ('Conclusion', 'EduTrack centralise le suivi des séances, l’émargement QR, la progression pédagogique et les honoraires.', BLUE),
    ('Recommandations', 'Prévoir un déploiement sécurisé : HTTPS, sauvegardes régulières et formation des utilisateurs.', GREEN),
    ('Perspectives', 'Étendre la solution avec un module étudiant, les notifications, la signature numérique et Mobile Money.', ORANGE),
]
for i, (title, body, accent) in enumerate(items, 1):
    card(conclusion, 0.42 + (i - 1) * 4.24, 2.00, 4.02, 3.94, i, title, body, accent)

# Remove the dedicated recommendations slide, retaining all other slides and the closing slide.
review_element = None
for element, slide in zip(list(prs.slides._sldIdLst), list(prs.slides)):
    if slide.part == recommendations.part:
        review_element = element
        break
if review_element is not None:
    prs.slides._sldIdLst.remove(review_element)

remaining = list(prs.slides)
reorder(prs, remaining)
renumber(prs)
prs.save(OUTPUT)
print(f'Created: {OUTPUT} ({len(prs.slides)} slides)')
