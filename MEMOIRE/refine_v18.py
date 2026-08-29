from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SOURCE = Path('La finalisima present_V15.before-final-structure.pptx')
OUTPUT = Path('La finalisima present_V18.pptx')
LOGO = Path('_edutrack_logo.png')
NAVY, BLUE, SKY = RGBColor(15, 66, 114), RGBColor(24, 104, 171), RGBColor(103, 181, 229)
BACKGROUND, WHITE = RGBColor(248, 250, 252), RGBColor(255, 255, 255)
SLATE, MUTED, BORDER = RGBColor(48, 65, 82), RGBColor(101, 122, 140), RGBColor(211, 226, 236)
GREEN, ORANGE = RGBColor(16, 135, 105), RGBColor(214, 121, 32)


def fill(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.color.rgb = color


def text(slide, value, x, y, w, h, size, color=SLATE, bold=False, align=PP_ALIGN.LEFT, font='Aptos', valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame; frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]; paragraph.alignment = align; paragraph.text = value
    for run in paragraph.runs:
        run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, color, rounded=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(shape, color)
    if rounded: shape.adjustments[0] = 0.10
    return shape


def circle(slide, x, y, d, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)); fill(shape, color); return shape


def clear(slide):
    tree = slide.shapes._spTree
    for shape in list(slide.shapes): tree.remove(shape._element)


def header(slide, title, subtitle):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BACKGROUND
    rect(slide, 0, 0, 13.333, 0.12, BLUE); rect(slide, 0.42, 0.48, 0.72, 0.06, BLUE)
    text(slide, title, 0.42, 0.68, 10, 0.56, 27, NAVY, True, font='Aptos Display')
    text(slide, subtitle, 0.42, 1.28, 9.7, 0.30, 11.5, MUTED)
    if LOGO.exists(): slide.shapes.add_picture(str(LOGO), Inches(10.88), Inches(0.47), width=Inches(1.55), height=Inches(0.45))
    rect(slide, 0.42, 7.12, 1.48, 0.045, SKY)


def card(slide, x, y, w, h, number, title, body, accent):
    shape = rect(slide, x, y, w, h, WHITE, True); shape.line.color.rgb = BORDER; shape.line.width = Pt(1)
    circle(slide, x + 0.28, y + 0.25, 0.46, accent)
    text(slide, str(number), x + 0.28, y + 0.25, 0.46, 0.46, 11, WHITE, True, PP_ALIGN.CENTER)
    text(slide, title, x + 0.92, y + 0.23, w - 1.18, 0.36, 15, NAVY, True)
    text(slide, body, x + 0.30, y + 0.78, w - 0.60, h - 0.96, 11.5, SLATE, False, valign=MSO_ANCHOR.TOP)


def arrange(prs, slides):
    mapping = {slide.part: element for slide, element in zip(prs.slides, prs.slides._sldIdLst)}
    target = [mapping[slide.part] for slide in slides]
    for element in list(prs.slides._sldIdLst): prs.slides._sldIdLst.remove(element)
    for element in target: prs.slides._sldIdLst.append(element)


def renumber(prs):
    total = len(prs.slides)
    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not getattr(shape, 'has_text_frame', False): continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if re.fullmatch(r'\d{1,2}\s*/\s*\d{1,2}', run.text.strip()): run.text = f'{index:02d} / {total:02d}'


prs = Presentation(SOURCE)
slides = list(prs.slides)
intro, recommendations, conclusion, closing = slides[1], slides[4], slides[12], slides[13]

clear(intro); header(intro, 'Introduction', 'Présentation du projet de fin de cycle')
text(intro, 'EduTrack est une application web et mobile de suivi pédagogique.', 0.70, 2.05, 11.95, 0.56, 25, NAVY, True, PP_ALIGN.CENTER, 'Aptos Display')
text(intro, 'Elle permet de gérer les séances, les emplois du temps, les émargements\npar QR code, les fiches de progression et les honoraires des enseignants.', 1.10, 2.98, 11.10, 0.82, 17, SLATE, align=PP_ALIGN.CENTER)
rect(intro, 2.14, 4.38, 9.05, 0.86, WHITE, True)
text(intro, 'Une solution unique pour faciliter le suivi des activités pédagogiques.', 2.45, 4.62, 8.45, 0.32, 16, BLUE, True, PP_ALIGN.CENTER)

clear(conclusion); header(conclusion, 'Conclusion', 'Bilan du projet EduTrack')
for i, (title, body, accent) in enumerate([
    ('Suivi centralisé', 'Les informations liées aux classes, séances et enseignants sont regroupées dans une même plateforme.', BLUE),
    ('Émargement traçable', 'Le QR code associe l’enseignant à la séance et améliore la fiabilité du suivi.', GREEN),
    ('Honoraires automatisés', 'Les séances émargées et validées servent de base au calcul des honoraires.', ORANGE),
], 1): card(conclusion, 0.42 + (i - 1) * 4.24, 2.00, 4.02, 3.94, i, title, body, accent)

clear(recommendations); header(recommendations, 'Recommandations & perspectives', 'Actions proposées pour accompagner l’évolution d’EduTrack')
for i, (title, body, accent) in enumerate([
    ('Déploiement sécurisé', 'Utiliser HTTPS, des sauvegardes régulières et une configuration de production maîtrisée.', BLUE),
    ('Accompagnement des utilisateurs', 'Former administrateurs et enseignants afin d’assurer une adoption durable de la plateforme.', GREEN),
    ('Évolutions fonctionnelles', 'Prévoir le module étudiant, les notifications, la signature numérique et le paiement Mobile Money.', ORANGE),
], 1): card(recommendations, 0.42 + (i - 1) * 4.24, 2.00, 4.02, 3.94, i, title, body, accent)

# Retain every original slide, but put recommendations directly after conclusion.
base = [slide for slide in slides if slide.part not in (recommendations.part, closing.part)]
conclusion_position = next(i for i, slide in enumerate(base) if slide.part == conclusion.part)
base.insert(conclusion_position + 1, recommendations)
base.append(closing)
arrange(prs, base)
renumber(prs)
prs.save(OUTPUT)
print(f'Created: {OUTPUT} ({len(prs.slides)} slides)')
