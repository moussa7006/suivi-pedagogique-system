from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SOURCE = Path('La finalisima present_V19.before-detailed-summary.pptx')
OUTPUT = Path('La finalisima present_V20.pptx')
LOGO = Path('_edutrack_logo.png')

NAVY = RGBColor(15, 66, 114)
BLUE = RGBColor(24, 104, 171)
SKY = RGBColor(103, 181, 229)
BACKGROUND = RGBColor(248, 250, 252)
WHITE = RGBColor(255, 255, 255)
SLATE = RGBColor(48, 65, 82)
MUTED = RGBColor(101, 122, 140)
BORDER = RGBColor(211, 226, 236)
GREEN = RGBColor(16, 135, 105)
ORANGE = RGBColor(214, 121, 32)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color


def add_text(slide, value, x, y, w, h, size, color=SLATE, bold=False,
             align=PP_ALIGN.LEFT, font='Aptos', valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.text = value
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, color, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    fill(shape, color)
    if rounded:
        shape.adjustments[0] = 0.10
    return shape


def clear(slide):
    tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        tree.remove(shape._element)


def header(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BACKGROUND
    rect(slide, 0, 0, 13.333, 0.12, BLUE)
    rect(slide, 0.42, 0.48, 0.72, 0.06, BLUE)
    add_text(slide, 'Conclusion, recommandations & perspectives', 0.42, 0.68, 10.2, 0.56, 25, NAVY, True, font='Aptos Display')
    add_text(slide, 'Bilan du projet, actions de déploiement et évolutions envisagées', 0.42, 1.28, 9.8, 0.30, 11.5, MUTED)
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(10.88), Inches(0.47), width=Inches(1.55), height=Inches(0.45))
    rect(slide, 0.42, 7.12, 1.48, 0.045, SKY)


def section_label(slide, value, x, y, color):
    rect(slide, x, y, 2.42, 0.38, color, rounded=True)
    add_text(slide, value, x, y + 0.03, 2.42, 0.28, 10, WHITE, True, PP_ALIGN.CENTER)


def card(slide, x, y, w, h, title, body, accent):
    shape = rect(slide, x, y, w, h, WHITE, rounded=True)
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(1)
    rect(slide, x, y, 0.08, h, accent)
    add_text(slide, title, x + 0.27, y + 0.18, w - 0.48, 0.28, 13, NAVY, True)
    add_text(slide, body, x + 0.27, y + 0.55, w - 0.48, h - 0.72, 10.2, SLATE, False, valign=MSO_ANCHOR.TOP)


prs = Presentation(SOURCE)
slide = prs.slides[11]
clear(slide)
header(slide)

section_label(slide, 'CONCLUSION', 0.42, 1.82, BLUE)
section_label(slide, 'RECOMMANDATIONS & PERSPECTIVES', 6.80, 1.82, GREEN)

# Detailed explanations retained from the previous conclusion and recommendations slides.
left_cards = [
    ('Suivi centralisé', 'Les informations liées aux classes, séances et enseignants sont regroupées dans une même plateforme.'),
    ('Émargement traçable', 'Le QR code associe l’enseignant à la séance et améliore la fiabilité du suivi.'),
    ('Honoraires automatisés', 'Les séances émargées et validées servent de base au calcul des honoraires.'),
]
right_cards = [
    ('Déploiement sécurisé', 'Utiliser HTTPS, des sauvegardes régulières et une configuration de production maîtrisée.'),
    ('Accompagnement des utilisateurs', 'Former administrateurs et enseignants afin d’assurer une adoption durable de la plateforme.'),
    ('Évolutions fonctionnelles', 'Prévoir le module étudiant, les notifications, la signature numérique et le paiement Mobile Money.'),
]
for index, (title, body) in enumerate(left_cards):
    card(slide, 0.42, 2.34 + index * 1.43, 5.94, 1.20, title, body, [BLUE, GREEN, ORANGE][index])
for index, (title, body) in enumerate(right_cards):
    card(slide, 6.80, 2.34 + index * 1.43, 6.10, 1.20, title, body, [GREEN, BLUE, ORANGE][index])

prs.save(OUTPUT)
print(f'Created: {OUTPUT}')
