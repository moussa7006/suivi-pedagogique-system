from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

FILE = 'La finalisima present_V15.pptx'
NAVY = RGBColor(15, 66, 114)
MUTED = RGBColor(101, 122, 140)

prs = Presentation(FILE)
replacements = {
    'MOUSSA BALLA KEITA &\nAYA BOURAMA': 'MOUSSA BALLA KEITA & AYA BOURAMA',
    'MOUSSA BALLA KEITA & /\nAYA BOURAMA': 'MOUSSA BALLA KEITA & AYA BOURAMA',
    'Tout au même endroit': 'Décisions ralenties',
    'Données vérifiées': 'Présences difficiles à prouver',
    'Calculs et rapports': 'Honoraires et rapports peu fiables',
    'Conclusion & perspectives': 'Conclusion, recommandations & perspectives',
}

for slide in prs.slides:
    for shape in slide.shapes:
        if not getattr(shape, 'has_text_frame', False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text in replacements:
                    run.text = replacements[run.text]

last = prs.slides[-1]
box = last.shapes.add_textbox(Inches(6.78), Inches(5.76), Inches(4.55), Inches(0.32))
tf = box.text_frame
tf.clear()
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
p.text = 'Questions & échanges — 30 minutes'
for run in p.runs:
    run.font.name = 'Aptos'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = NAVY

prs.save(FILE)
print(f'Finalized: {FILE}')
