from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SOURCE = Path('La finalisima present_V11.pptx')
LOGO = Path('_edutrack_logo.png')
HERO = Path('_cover_visual.png')

NAVY = RGBColor(15, 66, 114)
BLUE = RGBColor(24, 104, 171)
SKY = RGBColor(103, 181, 229)
PALE = RGBColor(232, 245, 253)
WHITE = RGBColor(255, 255, 255)
SLATE = RGBColor(48, 65, 82)
MUTED = RGBColor(101, 122, 140)

prs = Presentation(SOURCE)
cover = prs.slides[0]

# Reuse the exact branding assets embedded in the cover.
for shape, output in ((cover.shapes[3], LOGO), (cover.shapes[14], HERO)):
    output.write_bytes(shape.image.blob)


def rgb_fill(shape, color, transparency=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if transparency is not None:
        shape.fill.transparency = transparency
    shape.line.fill.background()


def add_text(slide, text, x, y, w, h, size, color, bold=False,
             font='Aptos', align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
             spacing=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = text
    if spacing is not None:
        p.space_after = Pt(spacing)
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_rect(slide, x, y, w, h, color, radius=False, transparency=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    rgb_fill(shape, color, transparency)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def add_oval(slide, x, y, w, h, color, transparency=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    rgb_fill(shape, color, transparency)
    return shape


# Update terminology to reflect the system's automatic validation.
replacements = {
    'Valider': 'Enregistrer automatiquement',
    'Trace fiable': 'Émargement tracé',
    'Validation traçable': 'Émargement traçable',
    'Calcul plus transparent': 'Calcul automatique à partir des séances émargées',
}
for slide in prs.slides:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                for before, after in replacements.items():
                    if before in run.text:
                        run.text = run.text.replace(before, after)

# Introduce a restrained shared identity on content slides (2 through 8).
for index in range(1, 8):
    slide = prs.slides[index]
    # Accent line and compact brand signature stay above the existing content.
    add_rect(slide, 0.38, 0.30, 0.68, 0.055, BLUE)
    add_text(slide, 'EduTrack', 10.55, 0.17, 1.05, 0.26, 11, NAVY, True, align=PP_ALIGN.RIGHT)
    add_text(slide, f'{index + 1:02d} / 09', 11.72, 0.17, 0.78, 0.26, 9, MUTED, False, align=PP_ALIGN.RIGHT)
    # Bottom marker establishes a consistent finish without covering slide content.
    add_rect(slide, 0.38, 7.18, 1.45, 0.04, SKY)

# Rebuild final slide with the same visual language as the title cover.
last = prs.slides[-1]
sp_tree = last.shapes._spTree
for shape in list(last.shapes):
    sp_tree.remove(shape._element)

# Cover-inspired background composition.
add_oval(last, -2.0, -2.0, 4.0, 4.0, PALE)
add_oval(last, 10.0, 4.1, 4.0, 4.0, NAVY)
add_rect(last, 0.37, 0.17, 1.62, 0.82, PALE, radius=True)
last.shapes.add_picture(str(LOGO), Inches(0.42), Inches(0.31), width=Inches(1.47), height=Inches(0.43))

# Left: simple, prominent closing message.
add_text(last, 'PROJET DE FIN DE CYCLE', 0.42, 1.45, 5.5, 0.28, 11, BLUE, True, spacing=0)
add_text(last, 'Merci pour\nvotre attention', 0.42, 1.90, 6.05, 1.45, 34, NAVY, True, font='Aptos Display')
add_rect(last, 0.42, 3.55, 1.05, 0.07, SKY)
add_text(last, 'EduTrack', 0.42, 3.86, 2.4, 0.45, 23, BLUE, True, font='Aptos Display')
add_text(last, 'Une solution web et mobile pour un suivi\npédagogique plus fiable et plus fluide.', 0.42, 4.43, 5.65, 0.72, 16, SLATE)

# Bottom presenter treatment echoes the cover slide.
add_rect(last, 0.0, 6.00, 4.40, 0.80, NAVY)
add_text(last, 'PRÉSENTÉ PAR', 0.54, 6.12, 1.6, 0.20, 8, SKY, True)
add_text(last, 'MOUSSA BALLA KEITA  ·  AYA BOURAMA', 0.54, 6.37, 3.45, 0.22, 10.5, WHITE, True)

# Right: visual panel and clean question prompt, matching the cover's asymmetric layout.
add_rect(last, 7.10, 0.35, 5.85, 3.52, PALE, radius=True)
last.shapes.add_picture(str(HERO), Inches(7.13), Inches(0.40), width=Inches(5.77), height=Inches(3.35))
add_oval(last, 10.85, 2.00, 2.35, 4.55, BLUE)
add_rect(last, 6.70, 4.35, 4.42, 1.38, WHITE, radius=True)
add_text(last, 'Questions ?', 7.03, 4.58, 3.72, 0.48, 29, NAVY, True, font='Aptos Display', align=PP_ALIGN.CENTER)
add_text(last, 'Nous serons ravis d’échanger avec vous.', 7.04, 5.15, 3.70, 0.25, 11.5, MUTED, False, align=PP_ALIGN.CENTER)

prs.save(SOURCE)
print('Presentation updated:', SOURCE)
