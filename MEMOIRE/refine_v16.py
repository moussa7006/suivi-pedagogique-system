from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SOURCE = Path('La finalisima present_V15.before-final-structure.pptx')
OUTPUT = Path('La finalisima present_V17.pptx')
LOGO = Path('_edutrack_logo.png')

NAVY = RGBColor(15, 66, 114)
BLUE = RGBColor(24, 104, 171)
SKY = RGBColor(103, 181, 229)
BACKGROUND = RGBColor(248, 250, 252)
WHITE = RGBColor(255, 255, 255)
SLATE = RGBColor(48, 65, 82)
MUTED = RGBColor(101, 122, 140)
BORDER = RGBColor(211, 226, 236)
GREEN = RGBColor(16, 135, 105)
ORANGE = RGBColor(214, 121, 32)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color


def add_text(slide, text, x, y, w, h, size, color=SLATE, bold=False,
             align=PP_ALIGN.LEFT, font='Aptos', valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = valign
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    paragraph.text = text
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_rect(slide, x, y, w, h, color, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(shape, color)
    if rounded:
        shape.adjustments[0] = 0.10
    return shape


def add_circle(slide, x, y, d, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    fill(shape, color)
    return shape


def clear_slide(slide):
    tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        tree.remove(shape._element)


def add_header(slide, title, subtitle):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BACKGROUND
    add_rect(slide, 0, 0, 13.333, 0.12, BLUE)
    add_rect(slide, 0.42, 0.48, 0.72, 0.06, BLUE)
    add_text(slide, title, 0.42, 0.68, 10, 0.56, 27, NAVY, True, font='Aptos Display')
    add_text(slide, subtitle, 0.42, 1.28, 9.6, 0.30, 11.5, MUTED)
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(10.88), Inches(0.47), width=Inches(1.55), height=Inches(0.45))
    add_rect(slide, 0.42, 7.12, 1.48, 0.045, SKY)


def add_card(slide, x, y, w, h, number, title, body, accent):
    card = add_rect(slide, x, y, w, h, WHITE, rounded=True)
    card.line.color.rgb = BORDER
    card.line.width = Pt(1)
    add_circle(slide, x + 0.28, y + 0.25, 0.46, accent)
    add_text(slide, str(number), x + 0.28, y + 0.25, 0.46, 0.46, 11, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, title, x + 0.92, y + 0.23, w - 1.18, 0.36, 15, NAVY, True)
    add_text(slide, body, x + 0.30, y + 0.78, w - 0.60, h - 0.96, 11.5, SLATE, False, valign=MSO_ANCHOR.TOP)


def add_recommendations_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    add_header(slide, 'Recommandations & perspectives', 'Actions proposées pour accompagner l’évolution d’EduTrack')
    cards = [
        ('Déploiement sécurisé', 'Utiliser HTTPS, des sauvegardes régulières et une configuration de production maîtrisée.'),
        ('Accompagnement des utilisateurs', 'Former administrateurs et enseignants afin d’assurer une adoption durable de la plateforme.'),
        ('Évolutions fonctionnelles', 'Prévoir le module étudiant, les notifications, la signature numérique et le paiement Mobile Money.'),
    ]
    positions = [(0.42, 2.00, 4.02, 3.94), (4.66, 2.00, 4.02, 3.94), (8.90, 2.00, 4.02, 3.94)]
    for index, ((title, body), position, accent) in enumerate(zip(cards, positions, [BLUE, GREEN, ORANGE]), 1):
        add_card(slide, *position, index, title, body, accent)
    return slide


def reorder(prs, slides):
    ids = {slide.part: element for slide, element in zip(prs.slides, prs.slides._sldIdLst)}
    elements = [ids[slide.part] for slide in slides]
    for element in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(element)
    for element in elements:
        prs.slides._sldIdLst.append(element)


def update_numbers(prs):
    total = len(prs.slides)
    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not getattr(shape, 'has_text_frame', False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if re.fullmatch(r'\d{1,2}\s*/\s*\d{1,2}', run.text.strip()):
                        run.text = f'{index:02d} / {total:02d}'


prs = Presentation(SOURCE)
slides = list(prs.slides)

# Rebuild the introduction in a simple, direct form.
intro = slides[1]
clear_slide(intro)
add_header(intro, 'Introduction', 'Présentation du projet de fin de cycle')
add_text(intro, 'EduTrack est une application web et mobile de suivi pédagogique.', 0.70, 2.05, 11.95, 0.56, 25, NAVY, True, align=PP_ALIGN.CENTER, font='Aptos Display')
add_text(intro, 'Elle permet de gérer les séances, les emplois du temps, les émargements\npar QR code, les fiches de progression et les honoraires des enseignants.', 1.10, 2.98, 11.10, 0.82, 17, SLATE, align=PP_ALIGN.CENTER)
add_rect(intro, 2.14, 4.38, 9.05, 0.86, WHITE, rounded=True)
add_text(intro, 'Une solution unique pour faciliter le suivi des activités pédagogiques.', 2.45, 4.62, 8.45, 0.32, 16, BLUE, True, align=PP_ALIGN.CENTER)

# Rebuild the conclusion only; recommendations move to their dedicated next slide.
conclusion = slides[12]
clear_slide(conclusion)
add_header(conclusion, 'Conclusion', 'Bilan du projet EduTrack')
for index, (title, body, accent) in enumerate([
    ('Suivi centralisé', 'Les informations liées aux classes, séances et enseignants sont regroupées dans une même plateforme.', BLUE),
    ('Émargement traçable', 'Le QR code associe l’enseignant à la séance et améliore la fiabilité du suivi.', GREEN),
    ('Honoraires automatisés', 'Les séances émargées et validées servent de base au calcul des honoraires.', ORANGE),
], 1):
    add_card(conclusion, 0.42 + (index - 1) * 4.24, 2.00, 4.02, 3.94, index, title, body, accent)

# Remove the review-critique slide (current slide 5).
review_slide = slides[4]
slide_id_list = prs.slides._sldIdLst
for element, slide in zip(list(slide_id_list), list(prs.slides)):
    if slide.part == review_slide.part:
        slide_id_list.remove(element)
        break

remaining = list(prs.slides)
recommendations = add_recommendations_slide(prs)
closing = remaining[-1]
conclusion = next(slide for slide in remaining if slide.part == conclusion.part)

# Place recommendations directly after conclusion and before the final thank-you slide.
ordered = [slide for slide in remaining if slide.part != closing.part and slide.part != recommendations.part]
conclusion_index = next(i for i, slide in enumerate(ordered) if slide.part == conclusion.part)
ordered.insert(conclusion_index + 1, recommendations)
ordered.append(closing)
reorder(prs, ordered)
update_numbers(prs)
prs.save(OUTPUT)
print(f'Created: {OUTPUT} ({len(prs.slides)} slides)')
