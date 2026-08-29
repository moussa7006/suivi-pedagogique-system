from pathlib import Path
from pptx import Presentation

PRESENTATION = Path('La finalisima present_V12.pptx')

slides_content = {
    0: {
        'CONCEPTION ET DÉVELOPPEMENT D\'UNE APPLICATION\nWEB ET MOBILE DE SUIVI PÉDAGOGIQUE ET DE CALCUL\nDES HONORAIRES DES ENSEIGNANTS': (
            "CONCEPTION ET DÉVELOPPEMENT D'UNE APPLICATION\n"
            "WEB ET MOBILE DE SUIVI PÉDAGOGIQUE ET DE CALCUL\n"
            "DES HONORAIRES DES ENSEIGNANTS\n\n"
            "Soutenance — 15 minutes de présentation · 30 minutes d’échanges"
        ),
    },
    1: {
        'Problématique & objectifs': 'Introduction, problématique & objectifs',
        'EduTrack — Projet de Licence 3': 'EduTrack — Introduction · 2 min',
        'Avant': 'Contexte et problématique',
        'Objectifs': 'Objectifs du projet',
        'Feuilles Excel': 'Outils dispersés',
        'Suivi dispersé': 'Suivi difficile',
        'Formulaires papier': 'Émargements papier',
        'Saisie manuelle': 'Erreurs et lenteur',
        'Données isolées': 'Données non centralisées',
        'Peu de traçabilité': 'Traçabilité limitée',
        'Centraliser': 'Centraliser',
        'Tout au même endroit': 'Réunir les données pédagogiques',
        'Fiabiliser': 'Fiabiliser',
        'Données vérifiées': 'Rendre le suivi vérifiable',
        'Automatiser': 'Automatiser',
        'Calculs et rapports': 'Émargements, honoraires et rapports',
    },
    2: {
        'Solution proposée : EduTrack': 'Revue critique & solution proposée',
        'EduTrack — Projet de Licence 3': 'EduTrack — Revue critique · 2 min',
        'Web + Mobile': 'Limites relevées',
        'Administration web': 'Réponse EduTrack',
        'Gestion, planning, suivi': 'Une plateforme unique web et mobile',
        'Émargement QR': 'QR code sécurisé',
        'Scan rapide sur mobile': 'Scan mobile, rapide et traçable',
        'Fiche de progression': 'Fiche de progression',
        'Séance et contenu réalisé': 'Contenu réellement traité',
        'Honoraires': 'Honoraires fiables',
        'Calcul automatique à partir des séances émargées': 'Calcul à partir des séances émargées',
    },
    3: {
        'Fonctionnement du système': 'Objectifs & hypothèses de travail',
        'EduTrack — Projet de Licence 3': 'EduTrack — Hypothèses · 1 min 30',
        'Planifier': 'H1 — Centraliser',
        'Créer la séance': 'Les données réunies améliorent le suivi',
        'QR Code': 'H2 — Fiabiliser',
        'Identifiant unique': 'Le QR code renforce la traçabilité',
        'Scanner': 'H3 — Simplifier',
        'Action mobile': 'Le mobile facilite l’émargement',
        'Enregistrer automatiquement': 'H4 — Automatiser',
        'Émargement tracé': 'Les séances émargées fiabilisent le calcul',
        'Calculer': 'Objectif final',
        'Honoraires': 'Décision et paiement mieux suivis',
    },
    4: {
        'Architecture technique': 'Démarche méthodologique & analytique',
        'EduTrack - Projet de Licence 3': 'EduTrack — Méthodologie · 2 min',
        'Modèle MVC simplifié du système EduTrack': 'Analyse des besoins, modélisation UML et architecture MVC',
        'ACTEURS': '1. Analyse',
        'Admin / Enseignant': 'Acteurs et besoins',
        'VUE': '2. Conception',
        'Web et mobile\nAngular + Ionic': 'Cas d’utilisation et interfaces',
        'CONTROLEUR': '3. Développement',
        'API REST\nSpring Boot\nJWT': 'API REST sécurisée',
        'MODELE': '4. Validation',
        'Données du système\nPostgreSQL': 'Tests et démonstration',
        'Réponse vers les interfaces': 'Itérations entre besoins, conception et réalisation',
    },
    5: {
        'Technologies & Outils utilisés': 'Méthodologie de réalisation',
        'EduTrack — Projet de Licence 3': 'EduTrack — Réalisation · 1 min 30',
        'BACKEND': 'Backend sécurisé',
        'FRONTEND ( WEB & MOBILE )': 'Interfaces web et mobile',
        'OUTILS & CONCEPTION': 'Outils de qualité',
    },
    6: {
        'DÉMONSTRATION': 'Résultats & discussion',
        'EduTrack - Projet de Licence 3': 'EduTrack — Résultats · 3 min',
        'Mobile enseignant et interface web administrateur pendant la démonstration.': (
            'Résultats obtenus : les parcours administrateur et enseignant sont opérationnels.'
        ),
        'Mobile enseignant': 'Résultat 1 — Application enseignant',
        'Connexion, planning et scan QR': 'Connexion, planning, scan QR et fiche de progression',
        'Interface web administrateur': 'Résultat 2 — Administration centralisée',
        'Tableau de bord, séances, honoraires et écran public de salle': (
            'Tableau de bord, séances, QR actif, suivi et honoraires'
        ),
        'Parcours : connexion -> planning -> scan QR -> fiche -> honoraires': (
            'Discussion : le QR code relie la séance, l’émargement, la progression et les honoraires.'
        ),
    },
    7: {
        'Conclusion & perspectives': 'Conclusion & recommandations',
        'EduTrack — Projet de Licence 3': 'EduTrack — Conclusion · 1 min',
        'Conclusion': 'Conclusion',
        'Perspectives': 'Recommandations et perspectives',
        'Module étudiant': 'Module étudiant',
        'Espace dédié': 'Étendre le suivi aux apprenants',
        'Notifications': 'Notifications',
        'Email / SMS': 'Informer rapidement les acteurs',
        'Paiement': 'Paiement sécurisé',
        'Mobile money': 'Intégrer Mobile Money',
        'Signature': 'Signature numérique',
        'Enseignants': 'Renforcer la preuve d’émargement',
    },
    8: {
        'Questions ?': 'Questions & échanges — 30 minutes',
        'Nous serons ravis d’échanger avec vous.': (
            'Merci pour votre attention. Nous sommes prêts à répondre à vos questions.'
        ),
    },
}

prs = Presentation(PRESENTATION)
for slide_index, replacements in slides_content.items():
    slide = prs.slides[slide_index]
    for shape in slide.shapes:
        if not getattr(shape, 'has_text_frame', False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text in replacements:
                    run.text = replacements[run.text]

prs.save(PRESENTATION)
print(f'Presentation adapted: {PRESENTATION}')
