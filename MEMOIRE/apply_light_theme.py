from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SOURCE = Path('La finalisima present_V11.pptx')

# Palette claire et lisible en projection.
BACKGROUND = RGBColor(248, 250, 252)
CARD = RGBColor(255, 255, 255)
NAVY = RGBColor(15, 66, 114)
SLATE = RGBColor(48, 65, 82)
MUTED = RGBColor(90, 110, 128)
BORDER = RGBColor(211, 226, 236)
CYAN = RGBColor(24, 104, 171)
PALE_BLUE = RGBColor(232, 245, 253)
DIVIDER = RGBColor(203, 213, 225)

DARK_BACKGROUNDS = {
    (11, 17, 33),     # 0B1121
    (7, 13, 31),      # 070D1F
    (12, 24, 52),     # 0C1834
    (15, 23, 42),     # 0F172A
}
DARK_CARDS = {
    (23, 32, 51),     # 172033
    (19, 29, 49),     # 131D31
}
DARK_DETAILS = {
    (51, 65, 85),     # 334155
    (2, 6, 23),       # 020617
    (30, 41, 59),     # 1E293B
}


def solid_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def shape_rgb(shape):
    try:
        rgb = shape.fill.fore_color.rgb
        return tuple(rgb) if rgb is not None else None
    except (AttributeError, TypeError):
        return None


def set_text_color(shape, color, size=None, bold=None):
    if not getattr(shape, 'has_text_frame', False):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = color
            if size is not None:
                run.font.size = Pt(size)
            if bold is not None:
                run.font.bold = bold


prs = Presentation(SOURCE)

# Slides 2 à 8 : passage à une interface claire avec texte très contrasté.
for slide_index in range(1, 8):
    slide = prs.slides[slide_index]

    for shape in slide.shapes:
        rgb = shape_rgb(shape)
        if rgb in DARK_BACKGROUNDS:
            solid_fill(shape, BACKGROUND)
            shape.line.fill.background()
        elif rgb in DARK_CARDS:
            solid_fill(shape, CARD)
            shape.line.color.rgb = BORDER
            shape.line.width = Pt(1)
        elif rgb in DARK_DETAILS:
            solid_fill(shape, PALE_BLUE if rgb == (30, 41, 59) else DIVIDER)
            shape.line.fill.background()

        # Tous les textes de contenu utilisent désormais un contraste adapté
        # à un fond clair. Les éléments blancs de logo restent des images.
        if getattr(shape, 'has_text_frame', False) and shape.text.strip():
            set_text_color(shape, SLATE)

    # Titre principal et repères d'identité.
    set_text_color(slide.shapes[3], NAVY, bold=True)
    set_text_color(slide.shapes[8], MUTED)
    set_text_color(slide.shapes[9], MUTED)

    # Les libellés clés des cartes deviennent bleu marine.
    for shape in slide.shapes:
        if getattr(shape, 'has_text_frame', False) and shape.text.strip():
            if shape.top > Inches(1.1) and shape.top < Inches(6.8):
                set_text_color(shape, NAVY if shape.height >= Inches(0.25) else SLATE)

# Slide 4 : enlever la consigne interne qui était visible au jury.
slide4 = prs.slides[3]
for shape in list(slide4.shapes):
    if getattr(shape, 'has_text_frame', False) and 'Le public écoute le scénario' in shape.text:
        slide4.shapes._spTree.remove(shape._element)

# Slide 5 : les étiquettes de technologie conservent une couleur d'accent,
# mais sont lisibles sur la nouvelle surface claire.
slide5 = prs.slides[4]
for shape in slide5.shapes:
    if shape_rgb(shape) in {(0, 188, 212), (16, 185, 129), (67, 97, 238), (245, 158, 11)}:
        # Laisse les accents colorés; assure un texte sombre si présent.
        if getattr(shape, 'has_text_frame', False):
            set_text_color(shape, NAVY, bold=True)

# Ajout de la numérotation manquante à la slide de clôture.
last = prs.slides[-1]
number = last.shapes.add_textbox(Inches(11.72), Inches(0.17), Inches(0.78), Inches(0.26))
tf = number.text_frame
tf.clear()
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.RIGHT
run = p.add_run()
run.text = '09 / 09'
run.font.name = 'Aptos'
run.font.size = Pt(9)
run.font.color.rgb = MUTED

prs.save(SOURCE)
print(f'Applied light theme: {SOURCE}')
